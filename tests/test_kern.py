"""Regressionstests für die Platzhalterersetzung.

Das ist das Stück, an dem ein Fehler still bleibt: ein Schreiben sieht fertig
aus und hat trotzdem den falschen Wert drin. Deshalb hier abgesichert.

    python3 -m unittest discover tests
"""

import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "werkzeuge"))

from kern import (  # noqa: E402
    berechne,
    blockschluessel,
    ersetze_in_docx,
    ersetze_in_markdown,
    platzhalter,
    sammle_werte,
    text_aus_docx,
    wert_aus_pfad,
)


class Platzhalter(unittest.TestCase):
    def test_findet_platzhalter(self):
        self.assertEqual(platzhalter("{{a}} und {{ b_2 }}"), {"a", "b_2"})

    def test_blockmarker_sind_keine_platzhalter(self):
        self.assertEqual(platzhalter("{{?abhilfe}}"), set())
        self.assertEqual(blockschluessel("{{?abhilfe}}\nText\n{{/abhilfe}}"), {"abhilfe"})

    def test_unbekannter_platzhalter_bleibt_stehen(self):
        # Sichtbar stehen lassen ist besser, als still eine Lücke zu erzeugen.
        self.assertEqual(ersetze_in_markdown("{{fehlt}}", {}), "{{fehlt}}")


class Fristen(unittest.TestCase):
    def test_werktage_ueberspringen_wochenende(self):
        # Freitag + 1 Werktag ist Montag
        self.assertEqual(berechne("heute + 1 werktage", {}, date(2026, 7, 24)), "27.07.2026")

    def test_kalendertage_zaehlen_durch(self):
        self.assertEqual(berechne("heute + 3 tage", {}, date(2026, 7, 24)), "27.07.2026")

    def test_frist_ab_anderem_feld(self):
        werte = {"vertragsdatum": "14.03.2026"}
        self.assertEqual(berechne("vertragsdatum + 10 werktage", werte), "27.03.2026")

    def test_fehlende_basis_faellt_auf(self):
        with self.assertRaises(ValueError):
            berechne("vertragsdatum + 5 tage", {})


class Werte(unittest.TestCase):
    def test_verschachtelter_pfad(self):
        self.assertEqual(wert_aus_pfad({"a": {"b": "c"}}, "a.b"), "c")
        self.assertIsNone(wert_aus_pfad({"a": {}}, "a.b"))

    def test_offene_felder_werden_gemeldet(self):
        meta = {
            "felder": [
                {"key": "ag", "quelle": "projekt", "pfad": "auftraggeber.firma"},
                {"key": "fehlt", "quelle": "projekt", "pfad": "gibt.es.nicht"},
                {"key": "frage", "quelle": "abfrage", "frage": "?"},
                {"key": "optional", "quelle": "abfrage", "frage": "?", "pflicht": False},
                {"key": "fix", "quelle": "fest", "wert": "Text"},
            ]
        }
        werte, offen = sammle_werte(meta, {"auftraggeber": {"firma": "Musterbau"}}, {})
        self.assertEqual(werte["ag"], "Musterbau")
        self.assertEqual(werte["fix"], "Text")
        self.assertEqual(sorted(offen), ["fehlt", "frage"])

    def test_vorgabe_schlaegt_projektwert(self):
        meta = {"felder": [{"key": "ag", "quelle": "projekt", "pfad": "x"}]}
        werte, offen = sammle_werte(meta, {"x": "aus Projekt"}, {"ag": "aus Abfrage"})
        self.assertEqual(werte["ag"], "aus Abfrage")
        self.assertEqual(offen, [])


class MarkdownErsetzung(unittest.TestCase):
    def test_leere_zeile_faellt_weg(self):
        text = "{{firma}}\n{{zusatz}}\n{{ort}}"
        self.assertEqual(
            ersetze_in_markdown(text, {"firma": "A", "zusatz": "", "ort": "B"}), "A\nB"
        )

    def test_optionaler_block_verschwindet_mit_ueberschrift(self):
        text = "vor\n{{?f}}\nÜberschrift\n{{f}}\n{{/f}}\nnach"
        self.assertEqual(ersetze_in_markdown(text, {"f": ""}), "vor\nnach")
        self.assertEqual(
            ersetze_in_markdown(text, {"f": "Inhalt"}), "vor\nÜberschrift\nInhalt\nnach"
        )


class DocxErsetzung(unittest.TestCase):
    """Die Fälle, die Word tatsächlich produziert."""

    def _vorlage(self, ordner: Path) -> Path:
        from docx import Document

        d = Document()
        d.sections[0].header.paragraphs[0].text = "Az. {{aktenzeichen}}"
        d.add_paragraph("{{firma}}")
        d.add_paragraph("{{zusatz}}")
        absatz = d.add_paragraph()
        for teil in ("Vertrag vom {{ver", "tragsdatum", "}}"):
            absatz.add_run(teil)
        tabelle = d.add_table(rows=1, cols=1)
        tabelle.rows[0].cells[0].text = "{{bauvorhaben}}"
        for zeile in ("{{?abhilfe}}", "Vorschlag", "{{abhilfe}}", "{{/abhilfe}}"):
            d.add_paragraph(zeile)
        pfad = ordner / "vorlage.docx"
        d.save(str(pfad))
        return pfad

    def _erzeuge(self, werte: dict) -> str:
        with tempfile.TemporaryDirectory() as tmp:
            ordner = Path(tmp)
            ziel = ordner / "ausgabe.docx"
            ersetze_in_docx(self._vorlage(ordner), ziel, werte)
            return text_aus_docx(ziel)

    def setUp(self):
        self.basis = {
            "aktenzeichen": "2451-DZ",
            "firma": "Musterbau GmbH",
            "zusatz": "",
            "vertragsdatum": "14.03.2026",
            "bauvorhaben": "Zentralküche",
            "abhilfe": "",
        }

    def test_kopfzeile_tabelle_und_zerrissener_platzhalter(self):
        # Word verteilt Platzhalter regelmäßig über mehrere Runs.
        ergebnis = self._erzeuge(self.basis)
        self.assertIn("Az. 2451-DZ", ergebnis)
        self.assertIn("Vertrag vom 14.03.2026", ergebnis)
        self.assertIn("Zentralküche", ergebnis)
        self.assertNotIn("{{", ergebnis)

    def test_leerer_adresszusatz_hinterlaesst_keine_luecke(self):
        # Der Absatz mit {{zusatz}} muss verschwinden, nicht leer stehen bleiben.
        ergebnis = self._erzeuge(self.basis)
        self.assertIn("Musterbau GmbH\nVertrag vom", ergebnis)

    def test_optionaler_block(self):
        self.assertNotIn("Vorschlag", self._erzeuge(self.basis))
        gefuellt = self._erzeuge({**self.basis, "abhilfe": "Zuleitung versetzen."})
        self.assertIn("Vorschlag", gefuellt)
        self.assertIn("Zuleitung versetzen.", gefuellt)

    def test_mehrzeiliger_wert_behaelt_umbrueche(self):
        ergebnis = self._erzeuge({**self.basis, "abhilfe": "Zeile eins\nZeile zwei"})
        self.assertIn("Zeile eins\nZeile zwei", ergebnis)


if __name__ == "__main__":
    unittest.main()
