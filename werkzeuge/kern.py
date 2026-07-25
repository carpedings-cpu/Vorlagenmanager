"""Gemeinsame Bausteine für Prüfung und Generierung.

Platzhalter haben die Form {{key}}. Sie funktionieren in Markdown-Vorlagen
und in DOCX-Vorlagen, dort auch in Tabellen sowie Kopf- und Fußzeilen.
"""

from __future__ import annotations

import re
from datetime import date, timedelta
from pathlib import Path

import yaml

PLATZHALTER = re.compile(r"\{\{\s*([a-z0-9_]+)\s*\}\}")
BLOCK_AUF = re.compile(r"^\s*\{\{\?\s*([a-z0-9_]+)\s*\}\}\s*$")
BLOCK_ZU = re.compile(r"^\s*\{\{/\s*([a-z0-9_]+)\s*\}\}\s*$")
QUELLEN = {"projekt", "vertrag", "abfrage", "berechnet", "fest"}
DATUMSFORMAT = "%d.%m.%Y"


# --- Laden -----------------------------------------------------------------


def lade_yaml(pfad: Path) -> dict:
    with pfad.open(encoding="utf-8") as f:
        return yaml.safe_load(f) or {}


def lade_vorlage(ordner: Path) -> dict:
    """Liest felder.yaml und ergänzt die aufgelösten Dateipfade."""
    meta = lade_yaml(ordner / "felder.yaml")
    meta["_ordner"] = ordner
    meta["_datei"] = ordner / meta.get("datei", "vorlage.docx")
    mail = meta.get("mail")
    meta["_mail"] = ordner / mail if mail else None
    return meta


def felder_nach_key(meta: dict) -> dict:
    return {f["key"]: f for f in meta.get("felder", [])}


# --- Platzhalter finden ----------------------------------------------------


def text_aus_docx(pfad: Path) -> str:
    """Gesamter Text einer DOCX inklusive Tabellen, Kopf- und Fußzeilen."""
    from docx import Document

    dok = Document(str(pfad))
    stuecke: list[str] = []

    def sammle(behaelter) -> None:
        for absatz in behaelter.paragraphs:
            stuecke.append(absatz.text)
        for tabelle in behaelter.tables:
            for zeile in tabelle.rows:
                for zelle in zeile.cells:
                    sammle(zelle)

    sammle(dok)
    for abschnitt in dok.sections:
        for teil in (abschnitt.header, abschnitt.footer, abschnitt.first_page_header):
            if teil is not None:
                sammle(teil)
    return "\n".join(stuecke)


def text_aus_datei(pfad: Path) -> str:
    if pfad.suffix.lower() == ".docx":
        return text_aus_docx(pfad)
    return pfad.read_text(encoding="utf-8")


def platzhalter(text: str) -> set[str]:
    return set(PLATZHALTER.findall(text))


def blockschluessel(text: str) -> set[str]:
    """Keys der optionalen Blöcke {{?key}} … {{/key}}."""
    keys = set()
    for zeile in text.split("\n"):
        for muster in (BLOCK_AUF, BLOCK_ZU):
            treffer = muster.match(zeile)
            if treffer:
                keys.add(treffer.group(1))
    return keys


# --- Werte ermitteln -------------------------------------------------------


def wert_aus_pfad(daten: dict, pfad: str):
    """Holt einen verschachtelten Wert, z. B. 'auftraggeber.firma'."""
    aktuell = daten
    for teil in pfad.split("."):
        if not isinstance(aktuell, dict) or teil not in aktuell:
            return None
        aktuell = aktuell[teil]
    return aktuell


def _parse_datum(wert: str) -> date:
    tag, monat, jahr = (int(t) for t in str(wert).strip().split("."))
    return date(jahr, monat, tag)


def _plus_werktage(start: date, anzahl: int) -> date:
    schritt = 1 if anzahl >= 0 else -1
    rest = abs(anzahl)
    aktuell = start
    while rest:
        aktuell += timedelta(days=schritt)
        if aktuell.weekday() < 5:
            rest -= 1
    return aktuell


FORMEL = re.compile(
    r"^\s*(?P<basis>heute|[a-z0-9_]+)"
    r"(?:\s*(?P<op>[+-])\s*(?P<anzahl>\d+)\s*(?P<einheit>werktage?|tage?))?\s*$"
)


def berechne(formel: str, werte: dict, heute: date | None = None) -> str:
    """Wertet eine Formel aus felder.yaml aus. Feiertage bleiben unberücksichtigt."""
    treffer = FORMEL.match(formel)
    if not treffer:
        raise ValueError(f"Formel nicht verstanden: {formel!r}")

    basis_name = treffer.group("basis")
    if basis_name == "heute":
        basis = heute or date.today()
    else:
        roh = werte.get(basis_name)
        if not roh:
            raise ValueError(f"Formel {formel!r} braucht das Feld {basis_name!r}")
        basis = _parse_datum(roh)

    if not treffer.group("op"):
        return basis.strftime(DATUMSFORMAT)

    anzahl = int(treffer.group("anzahl"))
    if treffer.group("op") == "-":
        anzahl = -anzahl
    if treffer.group("einheit").startswith("werktag"):
        ergebnis = _plus_werktage(basis, anzahl)
    else:
        ergebnis = basis + timedelta(days=anzahl)
    return ergebnis.strftime(DATUMSFORMAT)


def sammle_werte(meta: dict, projekt: dict, vorgaben: dict) -> tuple[dict, list[str]]:
    """Setzt die Werte aus Projektdaten, Vorgaben, Festwerten und Formeln zusammen.

    `vorgaben` enthält alles, was aus Abfrage und Vertragsdateien kommt und
    hat Vorrang. Zurück kommen die Werte und die Liste der Felder, die noch
    offen sind.
    """
    werte: dict[str, str] = {}
    offen: list[str] = []
    berechnete: list[dict] = []

    for feld in meta.get("felder", []):
        key = feld["key"]
        if key in vorgaben and vorgaben[key] not in (None, ""):
            werte[key] = str(vorgaben[key])
            continue

        quelle = feld.get("quelle")
        if quelle == "fest":
            werte[key] = str(feld.get("wert", ""))
        elif quelle == "projekt":
            wert = wert_aus_pfad(projekt, feld.get("pfad", key))
            if wert in (None, ""):
                if feld.get("pflicht", True):
                    offen.append(key)
                werte[key] = ""
            else:
                werte[key] = str(wert)
        elif quelle == "berechnet":
            berechnete.append(feld)
        else:  # abfrage, vertrag
            if feld.get("pflicht", True):
                offen.append(key)
            werte[key] = ""

    for feld in berechnete:
        werte[feld["key"]] = berechne(feld["formel"], werte)

    return werte, offen


# --- Ersetzen --------------------------------------------------------------


def ersetze(text: str, werte: dict) -> str:
    return PLATZHALTER.sub(lambda t: str(werte.get(t.group(1), t.group(0))), text)


def _nur_platzhalter(text: str) -> bool:
    """True, wenn die Zeile außer Platzhaltern nichts Sichtbares enthält."""
    return bool(text.strip()) and not PLATZHALTER.sub("", text).strip()


def _blockstatus(stapel: list[tuple[str, bool]]) -> bool:
    return all(behalten for _, behalten in stapel)


def _bloecke_markdown(text: str, werte: dict) -> str:
    """Entfernt optionale Blöcke, deren Feld leer geblieben ist."""
    zeilen: list[str] = []
    stapel: list[tuple[str, bool]] = []
    for zeile in text.split("\n"):
        auf = BLOCK_AUF.match(zeile)
        if auf:
            key = auf.group(1)
            stapel.append((key, bool(str(werte.get(key, "")).strip())))
            continue
        zu = BLOCK_ZU.match(zeile)
        if zu:
            if stapel and stapel[-1][0] == zu.group(1):
                stapel.pop()
            continue
        if _blockstatus(stapel):
            zeilen.append(zeile)
    return "\n".join(zeilen)


def _bloecke_docx(behaelter, werte: dict) -> None:
    entfernen = []
    stapel: list[tuple[str, bool]] = []
    for absatz in list(behaelter.paragraphs):
        auf = BLOCK_AUF.match(absatz.text)
        if auf:
            key = auf.group(1)
            stapel.append((key, bool(str(werte.get(key, "")).strip())))
            entfernen.append(absatz)
            continue
        zu = BLOCK_ZU.match(absatz.text)
        if zu:
            if stapel and stapel[-1][0] == zu.group(1):
                stapel.pop()
            entfernen.append(absatz)
            continue
        if not _blockstatus(stapel):
            entfernen.append(absatz)
    for absatz in entfernen:
        _absatz_entfernen(absatz)


def ersetze_in_markdown(text: str, werte: dict) -> str:
    zeilen = []
    for zeile in _bloecke_markdown(text, werte).split("\n"):
        neu = ersetze(zeile, werte)
        # Adresszusatz o. Ä. leer: Zeile fällt weg statt leer stehen zu bleiben
        if _nur_platzhalter(zeile) and not neu.strip():
            continue
        zeilen.append(neu)
    return "\n".join(zeilen)


def _setze_absatztext(absatz, text: str) -> None:
    """Schreibt den Text in den ersten Run, damit die Formatierung erhalten bleibt."""
    runs = absatz.runs
    if not runs:
        return
    for run in runs[1:]:
        run.text = ""
    erster = runs[0]
    teile = text.split("\n")
    erster.text = teile[0]
    for teil in teile[1:]:
        erster.add_break()
        erster.add_text(teil)


def _absatz_entfernen(absatz) -> None:
    absatz._element.getparent().remove(absatz._element)


def _ersetze_in_behaelter(behaelter, werte: dict) -> None:
    _bloecke_docx(behaelter, werte)
    for absatz in list(behaelter.paragraphs):
        # Runs zusammenfassen, damit auch über Runs verteilte Platzhalter greifen
        voll = "".join(run.text for run in absatz.runs)
        if not PLATZHALTER.search(voll):
            continue
        neu = ersetze(voll, werte)
        if _nur_platzhalter(voll) and not neu.strip():
            _absatz_entfernen(absatz)
            continue
        _setze_absatztext(absatz, neu)
    for tabelle in behaelter.tables:
        for zeile in tabelle.rows:
            for zelle in zeile.cells:
                _ersetze_in_behaelter(zelle, werte)


def ersetze_in_docx(quelle: Path, ziel: Path, werte: dict) -> None:
    from docx import Document

    dok = Document(str(quelle))
    _ersetze_in_behaelter(dok, werte)
    for abschnitt in dok.sections:
        for teil in (abschnitt.header, abschnitt.footer, abschnitt.first_page_header):
            if teil is not None:
                _ersetze_in_behaelter(teil, werte)
    ziel.parent.mkdir(parents=True, exist_ok=True)
    dok.save(str(ziel))
